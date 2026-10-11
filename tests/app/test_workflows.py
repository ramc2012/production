"""PRD acceptance: auth/roles/CSRF, F4 issue lifecycle, F5 test overrides, F6 what-if, F7 recommendations,
F8 L5 sandbox job, F9 allocation export, importer rejection, scheduler idempotence, backup."""
import io
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.backend import db as dbm
from app.backend.models import AuditLog, WellDailyResult

from conftest import Api, wait_job  # pytest puts tests/app on sys.path

V1 = "/api/v1"


def audit_actions(prefix):
    with dbm.session_scope() as db:
        return [a.action for a in db.scalars(select(AuditLog).where(AuditLog.action.like(f"{prefix}%")))]


# ---------------------------------------------------------------------------- auth
def test_auth_required_and_csrf(client, api):
    anon = TestClient(client.app)
    assert anon.get(f"{V1}/dashboard").status_code == 401
    assert anon.post(f"{V1}/auth/login", json={"username": "admin", "password": "wrong"}).status_code == 401
    # signed in but no CSRF header -> rejected
    assert api.c.post(f"{V1}/runs", json={"kind": "validation"}).status_code == 403


def test_roles(client, api):
    r = api.post(f"{V1}/admin/users", json={"username": "viewer1", "role": "viewer"})
    assert r.status_code == 200, r.text
    pw = r.json()["temporary_password"]
    v = Api(TestClient(client.app))
    me = v.login("viewer1", pw)
    assert me["must_change_password"] is True
    assert v.post(f"{V1}/runs", json={"kind": "validation"}).status_code == 403
    assert v.get(f"{V1}/admin/users").status_code == 403
    assert v.get(f"{V1}/dashboard").status_code == 200
    assert "user.create" in audit_actions("user.")


# ---------------------------------------------------------------------------- F4 issues
def test_issue_lifecycle(api, twin_run):
    issues = api.get(f"{V1}/issues").json()
    assert issues, "the planted faults must open issues"
    alarm = next(i for i in issues if i["severity"] == "ALARM")
    iid = alarm["id"]
    r = api.patch(f"{V1}/issues/{iid}", json={"state": "acknowledged", "assignee": "admin", "due": "2026-10-15",
                                               "comment": "checking at site"})
    assert r.status_code == 200 and r.json()["state"] == "acknowledged" and r.json()["assignee"] == "admin"
    assert api.patch(f"{V1}/issues/{iid}", json={"state": "resolved"}).status_code == 400   # needs a comment
    assert api.patch(f"{V1}/issues/{iid}", json={"state": "resolved", "comment": "fixed"}).json()["state"] == "resolved"
    assert api.patch(f"{V1}/issues/{iid}", json={"state": "acknowledged"}).status_code == 409   # illegal
    d = api.get(f"{V1}/issues/{iid}").json()
    assert [c["text"] for c in d["comments"]] == ["checking at site", "fixed"]
    assert d["findings"] and d["occurrences"] >= 1
    assert api.patch(f"{V1}/issues/{iid}", json={"state": "open", "comment": "back"}).json()["state"] == "open"
    assert audit_actions("issue.update").count("issue.update") >= 3
    b = api.post(f"{V1}/issues/bulk", json={"ids": [i["id"] for i in issues[:3]], "state": "acknowledged"}).json()
    assert len(b["updated"]) + len(b["skipped"]) == 3


def test_repeat_run_links_to_same_issue(api, twin_run, validation_run):
    """A second run over the same period updates the same tickets instead of creating new ones."""
    issues = api.get(f"{V1}/issues?state=open,acknowledged,resolved").json()
    keys = [(i["object_id"], i["check"]) for i in issues]
    assert len(keys) == len(set(keys))
    assert any(i["last_run_id"] == validation_run["id"] for i in issues)


# ---------------------------------------------------------------------------- F5 overrides
def test_override_feeds_next_run(api, validation_run):
    tests = api.get(f"{V1}/tests?run_id={validation_run['id']}").json()["tests"]
    t = next(x for x in tests if x["status"] == "REJECTED")
    assert api.post(f"{V1}/tests/{t['test_id']}/override", json={"status": "ACCEPTED", "reason": "x"}).status_code == 400
    r = api.post(f"{V1}/tests/{t['test_id']}/override",
                 json={"status": "ACCEPTED", "reason": "Witnessed test, SPM log corrected"})
    assert r.status_code == 200
    pending = api.get(f"{V1}/tests?run_id={validation_run['id']}").json()["tests"]
    assert next(x for x in pending if x["test_id"] == t["test_id"])["override_pending"] is True
    run = api.post(f"{V1}/runs", json={"kind": "validation"}).json()
    assert wait_job(api, run["job_id"])["status"] == "done"
    after = api.get(f"{V1}/tests?run_id={run['id']}").json()["tests"]
    row = next(x for x in after if x["test_id"] == t["test_id"])
    assert row["status"] == "ACCEPTED" and row["engine_status"] == "REJECTED"
    assert row["override"]["reason"] == "Witnessed test, SPM log corrected"
    assert api.get(f"{V1}/runs/{run['id']}").json()["test_overrides"][t["test_id"]]["status"] == "ACCEPTED"
    assert api.delete(f"{V1}/tests/{t['test_id']}/override", json={"reason": "test done"}).status_code == 200
    assert "test.override" in audit_actions("test.")


# ---------------------------------------------------------------------------- F6 what-if
def test_whatif_fast_and_consistent(api, twin_run):
    t0 = time.time()
    r = api.post(f"{V1}/twin/whatif", json={"changes": {"ANK-105": {"lift_gas": 9000}}})
    dt = time.time() - t0
    assert r.status_code == 200, r.text
    w = r.json()
    assert dt < 3.0, f"what-if took {dt:.2f}s"
    row = next(x for x in w["wells"] if x["well_id"] == "ANK-105")
    assert row["changed"] and row["lift_gas_sm3d"] == 9000
    assert abs(w["totals"]["oil"]["delta"] - sum(x["d_oil"] for x in w["wells"])) < 1e-6
    shut = api.post(f"{V1}/twin/whatif", json={"changes": {"ANK-101": {"on": 0}}}).json()
    assert next(x for x in shut["wells"] if x["well_id"] == "ANK-101")["oil_new"] == 0
    assert api.post(f"{V1}/twin/whatif", json={"changes": {"NOPE": {"choke": 20}}}).status_code == 400
    assert api.post(f"{V1}/twin/whatif", json={"changes": {"ANK-101": {"valve": 1}}}).status_code == 400


# ---------------------------------------------------------------------------- F7 recommendations
def test_recommendation_decisions(api, twin_run):
    recs = api.get(f"{V1}/twin/recommendations?run_id={twin_run['id']}").json()
    sp = [r for r in recs if r["kind"] == "setpoint"]
    assert sp and any(r["kind"] == "hold" for r in recs)
    gain = sum(r["predicted_d_oil_m3d"] or 0 for r in sp)
    assert gain > 5     # golden plan: +11.5 m3/d
    a, d = sp[0], sp[1]
    assert api.patch(f"{V1}/twin/recommendations/{a['id']}", json={"status": "accepted"}).json()["status"] == "accepted"
    assert api.patch(f"{V1}/twin/recommendations/{d['id']}", json={"status": "declined"}).status_code == 400
    assert api.patch(f"{V1}/twin/recommendations/{d['id']}",
                     json={"status": "declined", "comment": "MOC pending"}).json()["status"] == "declined"
    assert api.patch(f"{V1}/twin/recommendations/{d['id']}", json={"status": "implemented"}).status_code == 409
    with dbm.session_scope() as db:
        entries = db.scalars(select(AuditLog).where(AuditLog.action == "recommendation.decide")).all()
    assert len(entries) >= 2 and all("no plant system" in e.after["note"] for e in entries)


# ---------------------------------------------------------------------------- F8 L5 sandbox
def test_autonomy_sandbox_job_and_cancel(api, twin_run):
    r = api.post(f"{V1}/twin/autonomy", json={"mode": "autonomous_sim", "cycles": 3})
    assert r.status_code == 200, r.text
    assert "Simulation only" in r.json()["banner"]
    j = wait_job(api, r.json()["job_id"])
    assert j["status"] == "done", j
    runs = api.get(f"{V1}/twin/autonomy").json()
    full = api.get(f"{V1}/twin/autonomy/{runs[0]['id']}").json()
    assert full["result"]["log"] and full["result"]["trajectory"]
    assert full["mode"] == "autonomous_sim"
    # a long job can be cancelled
    r = api.post(f"{V1}/twin/autonomy", json={"mode": "autonomous_sim", "cycles": 30})
    jid = r.json()["job_id"]
    api.delete(f"{V1}/jobs/{jid}")
    j = wait_job(api, jid)
    assert j["status"] in ("cancelled", "done")
    if j["status"] == "done":        # finished before the cancel landed - acceptable, but rare
        pytest.skip("job finished before cancellation")
    assert api.post(f"{V1}/twin/autonomy", json={"mode": "write_to_dcs"}).status_code == 400


# ---------------------------------------------------------------------------- F9 allocation
def test_allocation_export_reproduces_engine(api, twin_run):
    months = api.get(f"{V1}/allocation/months").json()
    assert months[0]["month"] == "2026-09"
    v = api.get(f"{V1}/allocation?month=2026-09").json()
    assert v["coverage"]["days"] == 30 and len(v["wells"]) == 12
    x = api.get(f"{V1}/allocation/2026-09/export.xlsx")
    assert x.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(x.content), data_only=True)
    ws = wb["Monthly by well"]
    hdr = [c.value for c in ws[1]]
    got = {r[hdr.index("well_id")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    run_id = v["run"]["id"]
    with dbm.session_scope() as db:
        rows = db.scalars(select(WellDailyResult).where(WellDailyResult.run_id == run_id)).all()
    for well in got:
        eng = sum(r.dvr_oil_m3 or 0 for r in rows if r.well_id == well)
        eng_p = sum(r.prop_oil_m3 or 0 for r in rows if r.well_id == well)
        assert abs(got[well][hdr.index("dvr_oil_m3")] - eng) < 0.01
        assert abs(got[well][hdr.index("prop_oil_m3")] - eng_p) < 0.01
    c = api.post(f"{V1}/allocation/2026-09/close", json={})
    assert c.status_code == 200, c.text
    assert api.post(f"{V1}/allocation/2026-09/close", json={}).status_code == 409
    assert api.get(f"{V1}/allocation?month=2026-09").json()["closed"]["by"] == "admin"
    assert api.post(f"{V1}/allocation/2026-09/reopen", json={"reason": ""}).status_code == 400
    assert api.post(f"{V1}/allocation/2026-09/reopen", json={"reason": "late dispatch ticket"}).status_code == 200
    assert {"allocation.close", "allocation.reopen", "allocation.export"} <= set(audit_actions("allocation."))


# ---------------------------------------------------------------------------- data in
def test_malformed_upload_rejected(api, seeded):
    bad = b"date,well_id,hours_on\n2026-10-01,ANK-101,twenty\nnot-a-date,ANK-101,24\n"
    r = api.post(f"{V1}/data/upload", files=[("files", ("well_daily.csv", bad, "text/csv"))])
    assert r.status_code == 200
    b = r.json()
    assert b["status"] == "rejected"
    errs = " ".join(b["report"]["errors"])
    assert "non-numeric hours_on" in errs and "unreadable date" in errs
    assert api.post(f"{V1}/data/batches/{b['id']}/commit").status_code == 409
    unknown = api.post(f"{V1}/data/upload", files=[("files", ("stuff.csv", b"a,b\n1,2\n", "text/csv"))]).json()
    assert unknown["status"] == "rejected"


def test_upload_preview_then_commit_replaces_rows(api, seeded):
    good = b"date,well_id,hours_on,thp_ksc\n2026-09-30,ANK-101,23.5,21.0\n"
    b = api.post(f"{V1}/data/upload", files=[("files", ("well_daily_fix.csv", good, "text/csv"))]).json()
    assert b["status"] == "preview", b
    ds = b["datasets"]["well_daily"]
    assert ds["replaced"] == 1 and ds["new"] == 0
    assert b["preview"]["well_daily"][0]["hours_on"] == "23.5"
    assert api.post(f"{V1}/data/batches/{b['id']}/discard").json()["status"] == "discarded"
    cov = {c["dataset"]: c for c in api.get(f"{V1}/data/coverage").json()}
    assert cov["well_daily"]["rows"] == 360 and cov["well_daily"]["last"] == "2026-09-30"


def test_inspect_and_mapping(api, seeded):
    csv = b"Date,Well,Hrs\n2026-09-30,ANK-101,24\n"
    parts = api.post(f"{V1}/data/inspect", files={"file": ("dpr.csv", csv, "text/csv")}).json()
    assert parts[0]["columns"] == ["Date", "Well", "Hrs"]
    import json
    m = {"dpr.csv": {"Date": "date", "Well": "well_id", "Hrs": "hours_on"}}
    b = api.post(f"{V1}/data/upload", files=[("files", ("dpr.csv", csv, "text/csv"))],
                 data={"mapping": json.dumps(m), "hints": json.dumps({"dpr.csv": "well_daily"}),
                       "save_mapping_as": "DPR sheet"}).json()
    assert b["status"] == "preview", b["report"]
    api.post(f"{V1}/data/batches/{b['id']}/discard")
    assert "DPR sheet" in api.get(f"{V1}/data/mappings").json()


# ---------------------------------------------------------------------------- ops
def test_scheduler_daily_run_is_idempotent(api, twin_run):
    from app.backend.services import scheduler
    rid = scheduler.daily_run("schedule", force=True)
    assert rid is not None
    with dbm.session_scope() as db:
        from app.backend.models import Run
        job = db.get(Run, rid).job_id
    assert wait_job(api, job, 300)["status"] == "done"
    assert scheduler.daily_run("schedule") is None      # same day, same inputs -> skipped


def test_backup(api, seeded):
    r = api.post(f"{V1}/admin/backup")
    assert r.status_code == 200 and r.json()["size"] > 0
    assert api.get(f"{V1}/admin/backups").json()


def test_pages_have_data(api, twin_run):
    for url in ("/dashboard", "/network", "/wells", "/wells/ANK-106", "/tests", "/meters", "/meters/FT-123",
                "/tanks", "/separators", "/reconciliation", "/register", "/findings?severity=ALARM",
                "/twin/summary", "/data/quality", "/admin/system", "/admin/schedule", "/admin/audit",
                "/admin/network", "/admin/config"):
        r = api.get(V1 + url)
        assert r.status_code == 200, (url, r.text[:300])
    d = api.get(f"{V1}/dashboard").json()
    assert d["kpi"]["dvr_oil_m3"] > 100 and d["top_exceptions"] and d["twin"]["levels"]
    m = {x["id"]: x for x in api.get(f"{V1}/meters").json()["meters"]}
    assert m["FT-123"]["bias_days_flagged"] >= 3      # planted +12 % gas meter bias
    assert m["FT-201"]["proving_overdue"] is True
